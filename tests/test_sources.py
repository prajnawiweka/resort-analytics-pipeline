"""Source files contain exactly what each system would know on the business date."""

import csv
import json
import time
from datetime import date, datetime, timedelta

import pytest
from openpyxl import load_workbook

from ingestion.config import OPENING_DATE
from ingestion.simulator import build_world
from ingestion.sources import LAST_SAFE_DATE, PMS_RESERVATION_COLUMNS, write_business_date

BUSINESS_DATE = date(2026, 3, 10)


@pytest.fixture(scope="module")
def folder(tmp_path_factory):
    out = tmp_path_factory.mktemp("incoming")
    world = build_world()
    write_business_date(world, BUSINESS_DATE, out)
    write_business_date(world, date(2025, 12, 15), out)  # 2026 targets release
    return out


def read_reservations(folder, d):
    with (folder / d.isoformat() / "pms_reservations.csv").open() as f:
        return list(csv.DictReader(f))


def test_expected_files(folder):
    names = {p.name for p in (folder / BUSINESS_DATE.isoformat()).iterdir()}
    assert names == {
        "pms_reservations.csv",
        "pms_night_audit.csv",
        "ota_bookings.json",
        "reviews.json",
    }
    assert (folder / "2025-12-15" / "finance_targets_2026.xlsx").exists()


def test_reservations_never_come_from_the_future(folder):
    rows = read_reservations(folder, BUSINESS_DATE)
    assert rows and list(rows[0]) == PMS_RESERVATION_COLUMNS
    for r in rows:
        created = datetime.strptime(r["CREATED_AT"], "%d/%m/%Y %H:%M").date()
        assert created <= BUSINESS_DATE
        if r["STATUS"] == "CANCELLED":
            assert datetime.strptime(r["CANCELLED_AT"], "%d/%m/%Y %H:%M").date() == BUSINESS_DATE
        else:
            assert created == BUSINESS_DATE and r["CANCELLED_AT"] == ""


def test_ota_feed_points_at_pms_reservations(folder):
    pms = {r["RES_NO"] for r in read_reservations(folder, BUSINESS_DATE)}
    payload = json.loads((folder / BUSINESS_DATE.isoformat() / "ota_bookings.json").read_text())
    links = [b["hotel_reservation_no"] for b in payload["bookings"]]
    assert {x for x in links if x} <= pms
    assert all(b["booked_at"].endswith("Z") for b in payload["bookings"])


def test_review_scales(folder):
    reviews = json.loads((folder / BUSINESS_DATE.isoformat() / "reviews.json").read_text())
    for rv in reviews:
        scale = rv["rating"]["scale"]
        assert scale == (10 if rv["platform"] == "Booking.com" else 5)
        assert 1 <= rv["rating"]["overall"] <= scale


def test_finance_workbook_layout(folder):
    wb = load_workbook(folder / "2025-12-15" / "finance_targets_2026.xlsx")
    assert wb.sheetnames == ["Revenue Target", "Occupancy Target"]
    ws = wb["Revenue Target"]
    assert ws["A4"].value == "Property" and ws["B4"].value == "Jan-26"
    assert ws["A10"].value == "TOTAL"


@pytest.mark.parametrize("d", [date(2024, 12, 30), LAST_SAFE_DATE + timedelta(days=1)])
def test_out_of_range_dates_are_refused(tmp_path, d):
    with pytest.raises(ValueError):
        write_business_date(build_world(), d, tmp_path)


def test_opening_export_holds_all_prior_bookings(tmp_path):
    world = build_world()
    write_business_date(world, OPENING_DATE, tmp_path)
    rows = read_reservations(tmp_path, OPENING_DATE)
    unique = {r["RES_NO"] for r in rows}
    expected = {b.res_no for b in world.bookings if b.booked_at.date() <= OPENING_DATE}
    assert unique == expected


def test_files_are_byte_identical_when_regenerated(tmp_path):
    """The loader skips unchanged files by checksum, so regenerating must not change bytes."""
    world = build_world()
    first, second = tmp_path / "a", tmp_path / "b"
    write_business_date(world, date(2025, 12, 15), first)
    time.sleep(1.1)  # Excel metadata has one-second resolution
    write_business_date(world, date(2025, 12, 15), second)
    for path in (first / "2025-12-15").iterdir():
        assert path.read_bytes() == (second / "2025-12-15" / path.name).read_bytes(), path.name
