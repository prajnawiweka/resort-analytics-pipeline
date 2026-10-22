"""Renders the simulated world as the files each source system would send for one business date.

| File                         | System                | Format | Quirks the pipeline must handle       |
|------------------------------|-----------------------|--------|----------------------------------------|
| pms_reservations.csv         | Property Mgmt System  | CSV    | dd/mm/yyyy dates, hotel codes, ISO2     |
|                              |                       |        | countries, duplicate rows, stray spaces |
| pms_night_audit.csv          | Property Mgmt System  | CSV    | amounts as text with thousand separator |
| ota_bookings.json            | Channel manager       | JSON   | own property names, channel spellings,  |
|                              |                       |        | UTC timestamps, missing PMS link        |
| reviews.json                 | Review aggregator     | JSON   | 1-10 and 1-5 scales, missing sub-scores,|
|                              |                       |        | different booking references            |
| finance_targets_<year>.xlsx  | Finance (Excel)       | XLSX   | title rows, wide months, TOTAL row      |

All timestamps in the PMS are local Bali time (WITA, UTC+8); the JSON feeds use UTC.
"""

from __future__ import annotations

import csv
import json
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from ingestion.config import CH, COUNTRIES, HORIZON, OPENING_DATE, PROP, PROPERTIES, REPORT_START, RT
from ingestion.simulator import (
    Booking,
    World,
    booked_on,
    booked_until,
    cancelled_on,
    reviews_on,
    rng,
)

WITA_OFFSET = timedelta(hours=8)
LAST_SAFE_DATE = HORIZON - timedelta(days=184)  # later dates would miss far-ahead bookings
TARGET_RELEASES = {OPENING_DATE: 2025, date(2025, 12, 15): 2026, date(2026, 12, 15): 2027}

PMS_RESERVATION_COLUMNS = [
    "RES_NO", "HOTEL_CODE", "ROOM_TYPE", "ARRIVAL", "DEPARTURE", "NIGHTS", "ADULTS", "CHILDREN",
    "RATE", "TOTAL", "STATUS", "SOURCE", "GUEST_NAME", "GUEST_EMAIL", "GUEST_COUNTRY",
    "CREATED_AT", "CANCELLED_AT",
]  # fmt: skip
PMS_AUDIT_COLUMNS = [
    "AUDIT_DATE", "HOTEL_CODE", "ROOMS_AVAILABLE", "ROOMS_SOLD", "ROOM_REV", "FB_REV", "SPA_REV",
    "OTHER_REV",
]  # fmt: skip


def utc_iso(local: datetime) -> str:
    return (local - WITA_OFFSET).strftime("%Y-%m-%dT%H:%M:%SZ")


def pms_date(d: date) -> str:
    return d.strftime("%d/%m/%Y")


def pms_datetime(t: datetime) -> str:
    return t.strftime("%d/%m/%Y %H:%M")


def _records_for(world: World, d: date) -> list[Booking]:
    """Reservations the PMS exports for business date d: new or cancelled that day.
    The opening export instead holds everything booked up to and including that date."""
    if d == OPENING_DATE:
        return booked_until(world, d)
    seen: dict[str, Booking] = {}
    for b in booked_on(world, d) + cancelled_on(world, d):
        seen.setdefault(b.res_no, b)
    return list(seen.values())


def _status(b: Booking, d: date) -> str:
    return "CANCELLED" if b.cancelled_at and b.cancelled_at.date() <= d else "RESERVED"


def write_pms_reservations(world: World, d: date, path: Path) -> int:
    guests = {g.id: g for g in world.guests}
    r = rng("pms-dirt", d)
    rows = []
    for b in _records_for(world, d):
        g, ch = guests[b.guest_id], CH[b.channel_id]
        cancelled = _status(b, d) == "CANCELLED"
        email = g.email.upper() if r.random() < 0.05 else g.email
        source = ch.code + (" " if r.random() < 0.03 else "")
        rows.append([
            b.res_no, PROP[b.property_id].code, RT[b.room_type_id].code, pms_date(b.check_in),
            pms_date(b.check_out), b.nights, b.adults, b.children, b.rate, b.total,
            _status(b, d), source, f"{g.last_name.upper()}, {g.first_name}", email, g.country,
            pms_datetime(b.booked_at), pms_datetime(b.cancelled_at) if cancelled else "",
        ])  # fmt: skip
        if r.random() < 0.01:
            rows.append(list(rows[-1]))  # the PMS export occasionally repeats a line
    _write_csv(path, PMS_RESERVATION_COLUMNS, rows)
    return len(rows)


def write_pms_night_audit(world: World, d: date, path: Path) -> int:
    rows = []
    for p in PROPERTIES:
        a = world.audit[d, p.id]
        rows.append([
            pms_date(d), p.code, a.rooms_available, a.rooms_sold, f"{a.room_revenue:,.2f}",
            f"{a.fnb_revenue:,.2f}", f"{a.spa_revenue:,.2f}", f"{a.other_revenue:,.2f}",
        ])  # fmt: skip
    _write_csv(path, PMS_AUDIT_COLUMNS, rows)
    return len(rows)


def write_ota_bookings(world: World, d: date, path: Path) -> int:
    r = rng("ota-dirt", d)
    items = []
    for b in _records_for(world, d):
        if not b.ota_ref:
            continue
        ch, cancelled = CH[b.channel_id], _status(b, d) == "CANCELLED"
        commission = 0.0 if cancelled else round(b.total * ch.commission_pct / 100, 2)
        items.append({
            "ota_ref": b.ota_ref,
            "channel": r.choice(ch.ota_spellings),
            "property": PROP[b.property_id].ota_name,
            "hotel_reservation_no": None if b.ota_link_missing else b.res_no,
            "stay": {"check_in": b.check_in.isoformat(), "check_out": b.check_out.isoformat()},
            "total_amount": b.total,
            "commission": {"amount": commission, "currency": "USD"},
            "status": "cancelled" if cancelled else "confirmed",
            "booked_at": utc_iso(b.booked_at),
            "cancelled_at": utc_iso(b.cancelled_at) if cancelled else None,
        })  # fmt: skip
    payload = {"report_date": d.isoformat(), "property_group": "Tropical Haven", "bookings": items}
    path.write_text(json.dumps(payload, indent=1, ensure_ascii=False))
    return len(items)


def write_reviews(world: World, d: date, path: Path) -> int:
    by_res = {b.res_no: b for b in world.bookings}
    items = []
    for rv in reviews_on(world, d):
        b = by_res[rv.res_no]
        scale = 10 if rv.platform == "Booking.com" else 5
        factor = scale / 5
        if rv.platform in ("Google Reviews", "TripAdvisor"):
            scores = {"overall": rv.scores["overall"]}  # these platforms publish one score
        else:
            scores = {k: round(v * factor, 1) for k, v in rv.scores.items()}
        reference = {"Booking.com": b.ota_ref, "Post-stay Survey": b.res_no}.get(rv.platform)
        items.append({
            "review_id": rv.review_id,
            "platform": rv.platform,
            "property": PROP[rv.property_id].review_name,
            "booking_reference": reference,
            "reviewer_country": None if rv.platform == "Post-stay Survey"
            else COUNTRIES[rv.reviewer_country],
            "rating": {"scale": scale, **scores},
            "text": rv.text,
            "submitted_at": utc_iso(rv.submitted_at),
        })  # fmt: skip
    path.write_text(json.dumps(items, indent=1, ensure_ascii=False))
    return len(items)


def write_finance_targets(world: World, year: int, path: Path) -> int:
    wb = Workbook()
    sheets = [("Revenue Target", 0, "USD", "#,##0"), ("Occupancy Target", 1, "%", "0.0%")]
    for i, (title, idx, unit, fmt) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws["A1"] = f"Tropical Haven Resort Group - {title} {year} ({unit})"
        ws["A1"].font = Font(bold=True)
        ws["A2"] = "Approved by Finance. Do not edit."
        months = [date(year, m, 1).strftime("%b-%y") for m in range(1, 13)]
        ws.append([])
        ws.append(["Property", *months])
        for p in PROPERTIES:
            ws.append([p.name, *(world.targets[year, m, p.id][idx] for m in range(1, 13))])
        if idx == 0:  # finance adds a total line to the revenue sheet only
            first, last = 5, 4 + len(PROPERTIES)
            ws.append(["TOTAL", *(f"=SUM({c}{first}:{c}{last})" for c in "BCDEFGHIJKLM")])
        for row in ws.iter_rows(min_row=5, min_col=2):
            for cell in row:
                cell.number_format = fmt
    wb.save(path)
    return len(PROPERTIES)


def _write_csv(path: Path, header: list[str], rows: list[list]) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def write_business_date(world: World, d: date, out_root: Path) -> dict[str, int]:
    """Write every file the source systems send for business date d. Returns rows per file."""
    if not OPENING_DATE <= d <= LAST_SAFE_DATE:
        raise ValueError(f"business date must be between {OPENING_DATE} and {LAST_SAFE_DATE}")
    folder = out_root / d.isoformat()
    folder.mkdir(parents=True, exist_ok=True)
    written = {
        "pms_reservations.csv": write_pms_reservations(world, d, folder / "pms_reservations.csv"),
        "ota_bookings.json": write_ota_bookings(world, d, folder / "ota_bookings.json"),
    }
    if d >= REPORT_START:
        written["pms_night_audit.csv"] = write_pms_night_audit(
            world, d, folder / "pms_night_audit.csv"
        )
        written["reviews.json"] = write_reviews(world, d, folder / "reviews.json")
    if d in TARGET_RELEASES:
        year = TARGET_RELEASES[d]
        name = f"finance_targets_{year}.xlsx"
        written[name] = write_finance_targets(world, year, folder / name)
    return written
